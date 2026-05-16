"""Document loading utilities for TXT, PDF, and DOCX files."""

from pathlib import Path

from docx import Document
from pypdf import PdfReader

# ── OCR imports (remove these two lines to revert OCR support) ───────
import pytesseract
from pdf2image import convert_from_path
# ─────────────────────────────────────────────────────────────────────


# ████████████████████████████████████████████████████████████████████
# ██ OCR CONFIG
# ████████████████████████████████████████████████████████████████████

# To fully revert OCR: remove this section, the OCR FALLBACK section
# below, and the fallback block inside load_pdf_file.
pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Users\abinzaid.t\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"
)
_POPPLER_PATH = r"C:\poppler\poppler-26.02.0\Library\bin"


# ████████████████████████████████████████████████████████████████████
# ██ DISPATCHER
# ████████████████████████████████████████████████████████████████████

def load_document(file_path: Path) -> str:
    """Load text content from a supported document file."""
    extension = file_path.suffix.lower()

    if extension == ".txt":
        return load_txt_file(file_path)
    if extension == ".pdf":
        return load_pdf_file(file_path)
    if extension == ".docx":
        return load_docx_file(file_path)

    raise ValueError(f"Unsupported file type: '{extension}'. Supported: .txt, .pdf, .docx")


# ████████████████████████████████████████████████████████████████████
# ██ OCR FALLBACK
# ████████████████████████████████████████████████████████████████████

def _ocr_pdf(file_path: Path) -> str:
    """Convert PDF pages to images and extract text via Tesseract."""
    try:
        images = convert_from_path(
            file_path,
            poppler_path=_POPPLER_PATH,
            last_page=10,
        )
        pages = [pytesseract.image_to_string(img, lang="eng+ara") for img in images]
        return "\n".join(pages)
    except Exception as e:
        print(f"[loader] OCR failed: {e}")
        return ""


# ████████████████████████████████████████████████████████████████████
# ██ LOADERS
# ████████████████████████████████████████████████████████████████████

def load_txt_file(file_path: Path) -> str:
    """Load text content from a TXT file."""
    return file_path.read_text(encoding="utf-8")


def load_pdf_file(file_path: Path) -> str:
    """Extract text from a PDF file, with OCR fallback for scanned pages."""
    reader = PdfReader(file_path)
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages)

    # ── OCR fallback (remove this block to revert) ───────────────────
    if len(text.strip()) < 200:
        print(f"[loader] Scanned PDF detected — running OCR on '{file_path.name}'")
        ocr_text = _ocr_pdf(file_path)
        if ocr_text.strip():
            return ocr_text
        print(f"[loader] OCR returned empty — using original extracted text")
    # ── end OCR fallback ─────────────────────────────────────────────

    return text


def load_docx_file(file_path: Path) -> str:
    """Extract text content from a DOCX file."""
    doc = Document(file_path)
    paragraphs = [para.text for para in doc.paragraphs]
    return "\n".join(paragraphs)


# ████████████████████████████████████████████████████████████████████
# ██ DOCX HTML PREVIEW
# ████████████████████████████████████████████████████████████████████

_DOCX_PREVIEW_CSS = (
    "body{font-family:Georgia,'Times New Roman',serif;max-width:760px;"
    "margin:2rem auto;padding:0 2rem 3rem;color:#1e293b;line-height:1.85;font-size:0.95rem;}"
    "h1{font-size:1.7rem;font-weight:700;margin:1.5rem 0 0.5rem;color:#0f172a;}"
    "h2{font-size:1.3rem;font-weight:700;margin:1.25rem 0 0.4rem;color:#1e293b;}"
    "h3{font-size:1.05rem;font-weight:600;margin:1rem 0 0.3rem;}"
    "h4{font-size:0.95rem;font-weight:600;margin:0.75rem 0 0.25rem;}"
    "p{margin:0.45rem 0;}"
    "table{border-collapse:collapse;width:100%;margin:1rem 0;font-size:0.9rem;}"
    "td,th{border:1px solid #cbd5e1;padding:0.5rem 0.75rem;vertical-align:top;}"
    "th{background:#f1f5f9;font-weight:600;}"
    "tr:nth-child(even) td{background:#f8fafc;}"
)


def _html_escape(text: str) -> str:
    """Escape characters that have special meaning in HTML."""
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _runs_to_html(runs) -> str:
    """Convert paragraph runs to inline HTML preserving bold, italic, and underline."""
    parts = []
    for run in runs:
        if not run.text:
            continue
        text = _html_escape(run.text).replace("\n", "<br>")
        if run.bold:
            text = f"<strong>{text}</strong>"
        if run.italic:
            text = f"<em>{text}</em>"
        if run.underline:
            text = f"<u>{text}</u>"
        parts.append(text)
    return "".join(parts)


def _para_to_html(para) -> str:
    """Convert a single DOCX paragraph to an HTML element string."""
    inline = _runs_to_html(para.runs)
    if not inline.strip():
        return ""
    style_name = (para.style.name or "").lower() if para.style else ""
    arabic_ratio = sum(1 for c in para.text if "؀" <= c <= "ۿ") / max(len(para.text.strip()), 1)
    dir_attr = ' dir="rtl" style="text-align:right;"' if arabic_ratio > 0.15 else ""
    if "heading 1" in style_name:
        return f"<h1{dir_attr}>{inline}</h1>"
    if "heading 2" in style_name:
        return f"<h2{dir_attr}>{inline}</h2>"
    if "heading 3" in style_name:
        return f"<h3{dir_attr}>{inline}</h3>"
    if "heading" in style_name:
        return f"<h4{dir_attr}>{inline}</h4>"
    if "list" in style_name or "bullet" in style_name:
        indent = ' style="margin-left:1.5rem;"'
        return f"<p{indent}>• {inline}</p>"
    return f"<p{dir_attr}>{inline}</p>"


def _table_to_html(table) -> str:
    """Convert a DOCX table to an HTML table string."""
    rows = []
    for i, row in enumerate(table.rows):
        cells = []
        for cell in row.cells:
            content = _html_escape(cell.text)
            tag = "th" if i == 0 else "td"
            cells.append(f"<{tag}>{content}</{tag}>")
        rows.append(f"<tr>{''.join(cells)}</tr>")
    return f"<table>{''.join(rows)}</table>"


def load_docx_as_html(file_path: Path) -> str:
    """Convert a DOCX file to styled HTML for inline browser preview."""
    doc = Document(file_path)
    para_map = {id(p._element): p for p in doc.paragraphs}
    table_map = {id(t._element): t for t in doc.tables}
    body_parts = []
    for child in doc.element.body:
        child_id = id(child)
        if child_id in para_map:
            html = _para_to_html(para_map[child_id])
            if html:
                body_parts.append(html)
        elif child_id in table_map:
            body_parts.append(_table_to_html(table_map[child_id]))
    return (
        f"<!DOCTYPE html><html>"
        f"<head><meta charset=\"utf-8\"><style>{_DOCX_PREVIEW_CSS}</style></head>"
        f"<body>{''.join(body_parts)}</body>"
        f"</html>"
    )
